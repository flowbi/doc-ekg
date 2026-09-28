"""Document readers with stable source identifiers and bounded chunks."""
from dataclasses import dataclass
from hashlib import sha256
from html.parser import HTMLParser
from pathlib import Path
import shutil
import subprocess
import tempfile


PLAIN_TEXT_SUFFIXES = frozenset({".txt", ".text", ".md", ".markdown", ".mdown", ".rst",
                                 ".adoc", ".asciidoc", ".org"})
SUPPORTED_SUFFIXES = PLAIN_TEXT_SUFFIXES | frozenset({
    ".html", ".htm", ".rtf", ".pdf", ".doc", ".docx",
    ".ppt", ".pptx",
})


@dataclass(frozen=True)
class Chunk:
    source: str
    source_hash: str
    index: int
    text: str


def _word_text(path: Path) -> str:
    from docx import Document
    from docx.table import Table

    parts = []
    for block in Document(str(path)).iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        elif block.text.strip():
            parts.append(block.text.strip())
    return "\n".join(parts)


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript"}:
            self.hidden += 1
        elif not self.hidden and tag in {"p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript"}:
            self.hidden = max(0, self.hidden - 1)
        elif not self.hidden and tag in {"p", "div", "li", "tr", "td", "th"}:
            self.parts.append("\n" if tag != "td" and tag != "th" else " | ")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def _html_text(path: Path) -> str:
    parser = _HTMLText()
    parser.feed(path.read_text(encoding="utf-8-sig"))
    return "\n".join(line.strip() for line in "".join(parser.parts).splitlines() if line.strip())


def _slide_shapes(shapes):
    for shape in shapes:
        if shape.has_text_frame and shape.text.strip():
            yield shape.text.strip()
        if shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
                if any(cells):
                    yield " | ".join(cells)
        if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            yield from _slide_shapes(shape.shapes)


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts.append(f"Slide: {number}")
        parts.extend(_slide_shapes(slide.shapes))
    return "\n".join(parts)


def _convert_legacy(path: Path, target_suffix: str, libreoffice_bin: str | None) -> str:
    binary = libreoffice_bin or shutil.which("libreoffice") or shutil.which("soffice")
    if not binary:
        mac_binary = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
        binary = str(mac_binary) if mac_binary.is_file() else None
    if not binary:
        raise RuntimeError(f"Reading {path.suffix} requires LibreOffice; set --libreoffice")
    with tempfile.TemporaryDirectory(prefix="document-ekg-") as folder:
        temporary = Path(folder)
        profile = (temporary / "profile").as_uri()
        completed = subprocess.run(
            [binary, f"-env:UserInstallation={profile}", "--headless", "--convert-to",
             target_suffix.lstrip("."), "--outdir", str(temporary), str(path)],
            capture_output=True, text=True, timeout=120, check=False,
        )
        converted = temporary / (path.stem + target_suffix)
        if completed.returncode != 0 or not converted.is_file():
            raise RuntimeError(f"LibreOffice could not convert {path.name}: "
                               f"{completed.stderr.strip() or completed.stdout.strip()}")
        return _word_text(converted) if target_suffix == ".docx" else _pptx_text(converted)


def read_document(path: Path, libreoffice_bin: str | None = None) -> str:
    if path.suffix.lower() in PLAIN_TEXT_SUFFIXES:
        return path.read_text(encoding="utf-8-sig")
    if path.suffix.lower() in {".html", ".htm"}:
        return _html_text(path)
    if path.suffix.lower() == ".rtf":
        return _convert_legacy(path, ".docx", libreoffice_bin)
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if path.suffix.lower() == ".docx":
        return _word_text(path)
    if path.suffix.lower() == ".doc":
        return _convert_legacy(path, ".docx", libreoffice_bin)
    if path.suffix.lower() == ".pptx":
        return _pptx_text(path)
    if path.suffix.lower() == ".ppt":
        return _convert_legacy(path, ".pptx", libreoffice_bin)
    raise ValueError(f"Unsupported document type: {path}")


def iter_files(paths: list[Path]):
    found = set()
    for path in paths:
        candidates = path.rglob("*") if path.is_dir() else [path]
        for candidate in candidates:
            if (candidate.is_file() and not candidate.name.startswith("~$")
                    and candidate.suffix.lower() in SUPPORTED_SUFFIXES):
                resolved = candidate.resolve()
                if resolved not in found:
                    found.add(resolved)
                    yield resolved


def chunks(path: Path, max_chars: int = 5000, overlap: int = 300,
           libreoffice_bin: str | None = None) -> list[Chunk]:
    if max_chars <= 0 or not 0 <= overlap < max_chars:
        raise ValueError("Require max_chars > overlap >= 0")
    source = path.resolve().as_uri()
    content_hash = sha256(path.read_bytes()).hexdigest()
    text = read_document(path, libreoffice_bin)
    result = []
    step = max_chars - overlap
    for index, start in enumerate(range(0, len(text), step)):
        part = text[start:start + max_chars]
        if not part.strip():
            continue
        result.append(Chunk(source, content_hash, index, part))
        if start + max_chars >= len(text):
            break
    return result
